from datetime import date
from decimal import Decimal
import openpyxl
from django.test import TestCase
from apps.academic_core.models import PeriodoLectivo, Nivel, Clase, Seccion, Asignatura, Materia, TipoNota
from apps.people.models import Estudiante, EstudianteClase, Docente
from apps.accounts.models import User
from apps.reporting.models import ConfiguracionInstitucion
from apps.reporting.services import (
    generar_boletin_xlsx,
    generar_listado_estudiantes_xlsx,
    generar_resumen_docentes_xlsx
)


class ReportingServiceTests(TestCase):
    def setUp(self):
        self.institucion = ConfiguracionInstitucion.get_solo()
        self.periodo = PeriodoLectivo.objects.create(
            nombre="Año Lectivo Reporting 2026",
            fecha_inicio=date(2026, 1, 15),
            fecha_fin=date(2026, 11, 30),
            activo=True
        )
        self.nivel = Nivel.objects.create(nombre="Primaria", codigo="PRIM_R")
        self.clase = Clase.objects.create(nombre="4to Grado", periodo=self.periodo, nivel=self.nivel)
        self.seccion = Seccion.objects.create(nombre="Sección A", codigo="4A", clase=self.clase)

        self.tipo_bloque = TipoNota.objects.create(
            nombre="Bloque I", periodo=self.periodo, es_consolidado=True, orden=1
        )

        self.estudiante = Estudiante.objects.create(
            codigo_estudiante="EST-REP-01",
            primer_nombre="Gabriel",
            primer_apellido="Solis",
            genero="M"
        )
        EstudianteClase.objects.create(
            estudiante=self.estudiante, clase=self.clase, seccion=self.seccion, periodo=self.periodo
        )

        self.user_docente = User.objects.create_user(username="prof_solis", first_name="Gabriel", last_name="Solis")
        self.docente = Docente.objects.create(usuario=self.user_docente, codigo_empleado="DOC-001")

    def test_generar_boletin_xlsx_valid_excel(self):
        buffer = generar_boletin_xlsx(self.estudiante, self.periodo)
        self.assertIsNotNone(buffer)
        
        # Verificar que openpyxl puede cargar el libro de trabajo desde el buffer de memoria
        wb = openpyxl.load_workbook(buffer)
        self.assertIn("Boletín de Calificaciones", wb.sheetnames)
        ws = wb["Boletín de Calificaciones"]
        self.assertIn(self.institucion.nombre_institucion, str(ws.cell(row=1, column=1).value))

    def test_generar_listado_estudiantes_xlsx_valid_excel(self):
        buffer = generar_listado_estudiantes_xlsx(self.clase, self.seccion)
        wb = openpyxl.load_workbook(buffer)
        self.assertIn("Listado de Alumnos", wb.sheetnames)
        ws = wb["Listado de Alumnos"]
        self.assertEqual(ws.cell(row=5, column=2).value, "EST-REP-01")

    def test_generar_resumen_docentes_xlsx_valid_excel(self):
        buffer = generar_resumen_docentes_xlsx(self.periodo)
        wb = openpyxl.load_workbook(buffer)
        self.assertIn("Resumen de Docentes", wb.sheetnames)
        ws = wb["Resumen de Docentes"]
        self.assertEqual(ws.cell(row=5, column=2).value, "DOC-001")
