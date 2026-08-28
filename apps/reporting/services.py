from io import BytesIO
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from apps.academic_core.models import Materia, TipoNota, Clase
from apps.people.models import EstudianteClase, Docente
from apps.grading.models import Calificacion, ResumenAcademicoEstudiante
from .models import ConfiguracionInstitucion


import os


def _aplicar_estilos_encabezado(ws, title, institucion):
    """ Helper para formatear membrete del colegio en los reportes """
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    cell_nombre = ws.cell(row=1, column=1, value=institucion.nombre_institucion)
    cell_nombre.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    cell_nombre.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
    cell_titulo = ws.cell(row=2, column=1, value=title.upper())
    cell_titulo.font = Font(name="Calibri", size=12, bold=True, color="262626")
    cell_titulo.alignment = Alignment(horizontal="center", vertical="center")

    if institucion.logo and hasattr(institucion.logo, "path") and os.path.exists(institucion.logo.path):
        try:
            from openpyxl.drawing.image import Image as OpenpyxlImage
            img = OpenpyxlImage(institucion.logo.path)
            img.width = 60
            img.height = 40
            ws.add_image(img, "A1")
        except Exception:
            pass

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 20


def _autoajustar_columnas(ws):
    """ Helper para ajustar ancho de columnas según contenido """
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def generar_boletin_xlsx(estudiante, periodo):
    """
    Genera el boletín de calificaciones del estudiante en formato XLSX.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Boletín de Calificaciones"

    institucion = ConfiguracionInstitucion.get_solo()
    _aplicar_estilos_encabezado(ws, f"Boletín Oficial - {periodo.nombre}", institucion)

    ws.cell(row=4, column=1, value="Estudiante:").font = Font(bold=True)
    ws.cell(row=4, column=2, value=estudiante.get_full_name())

    ws.cell(row=4, column=5, value="Carnet:").font = Font(bold=True)
    ws.cell(row=4, column=6, value=estudiante.codigo_estudiante)

    inscripcion = EstudianteClase.objects.filter(estudiante=estudiante, periodo=periodo).first()
    clase_str = str(inscripcion.seccion) if inscripcion else "Sin Matrícula"
    ws.cell(row=5, column=1, value="Grado/Sección:").font = Font(bold=True)
    ws.cell(row=5, column=2, value=clase_str)

    headers = ["Materia", "Docente"]
    tipos_bloque = list(periodo.tipos_notas.filter(es_consolidado=True).order_by("orden"))
    for tb in tipos_bloque:
        headers.append(tb.nombre)
    headers.append("Promedio Final")

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    row_idx = 7
    for col_idx, text in enumerate(headers, 1):
        c = ws.cell(row=row_idx, column=col_idx, value=text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    materias = Materia.objects.filter(clase__periodo=periodo)
    if inscripcion:
        materias = materias.filter(clase=inscripcion.clase)

    row_idx += 1
    for mat in materias:
        docente_str = mat.docente.get_full_name() if mat.docente else "Sin Asignar"
        ws.cell(row=row_idx, column=1, value=mat.nombre).border = thin_border
        ws.cell(row=row_idx, column=2, value=docente_str).border = thin_border

        col_pos = 3
        for tb in tipos_bloque:
            calif = Calificacion.objects.filter(
                materia=mat, tipo_nota=tb, estudiante=estudiante, es_consolidado=True
            ).first()
            val = calif.resultado if (calif and calif.resultado is not None) else "-"
            c = ws.cell(row=row_idx, column=col_pos, value=val)
            c.border = thin_border
            c.alignment = Alignment(horizontal="center")
            col_pos += 1

        calif_final = Calificacion.objects.filter(
            materia=mat, estudiante=estudiante, es_consolidado=True, resultado__isnull=False
        )
        if calif_final.exists():
            prom = sum(c.resultado for c in calif_final) / Decimal(calif_final.count())
            c_prom = ws.cell(row=row_idx, column=col_pos, value=float(prom.quantize(Decimal("0.01"))))
        else:
            c_prom = ws.cell(row=row_idx, column=col_pos, value="-")
        c_prom.border = thin_border
        c_prom.font = Font(bold=True)
        c_prom.alignment = Alignment(horizontal="center")
        row_idx += 1

    row_idx += 1
    resumen = ResumenAcademicoEstudiante.objects.filter(estudiante=estudiante, periodo=periodo).first()
    prom_general = float(resumen.promedio_general) if (resumen and resumen.promedio_general) else "-"
    
    ws.cell(row=row_idx, column=1, value="PROMEDIO GENERAL ACUMULADO:").font = Font(bold=True)
    c_final = ws.cell(row=row_idx, column=2, value=prom_general)
    c_final.font = Font(bold=True, size=12, color="1F4E79")

    row_idx += 2
    ws.cell(row=row_idx, column=1, value=institucion.pie_de_pagina_boletin).font = Font(italic=True, size=9, color="595959")

    _autoajustar_columnas(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_listado_estudiantes_xlsx(clase, seccion=None):
    """
    Genera el listado oficial de estudiantes matriculados en una Clase/Sección en XLSX.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Listado de Alumnos"

    institucion = ConfiguracionInstitucion.get_solo()
    sub_titulo = f"Listado Oficial - {clase.nombre}"
    if seccion:
        sub_titulo += f" ({seccion.nombre})"

    _aplicar_estilos_encabezado(ws, sub_titulo, institucion)

    headers = ["No.", "Carnet", "Primer Nombre", "Segundo Nombre", "Primer Apellido", "Segundo Apellido", "Género", "Cédula", "Estado"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx, text in enumerate(headers, 1):
        c = ws.cell(row=4, column=col_idx, value=text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    inscripciones = EstudianteClase.objects.filter(clase=clase).select_related("estudiante")
    if seccion:
        inscripciones = inscripciones.filter(seccion=seccion)

    row_idx = 5
    for idx, inscrip in enumerate(inscripciones, 1):
        est = inscrip.estudiante
        ws.cell(row=row_idx, column=1, value=idx)
        ws.cell(row=row_idx, column=2, value=est.codigo_estudiante)
        ws.cell(row=row_idx, column=3, value=est.primer_nombre)
        ws.cell(row=row_idx, column=4, value=est.segundo_nombre)
        ws.cell(row=row_idx, column=5, value=est.primer_apellido)
        ws.cell(row=row_idx, column=6, value=est.segundo_apellido)
        ws.cell(row=row_idx, column=7, value=est.get_genero_display())
        ws.cell(row=row_idx, column=8, value=est.cedula_identidad or "-")
        ws.cell(row=row_idx, column=9, value=inscrip.get_estado_display())
        row_idx += 1

    _autoajustar_columnas(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_resumen_docentes_xlsx(periodo):
    """
    Genera el reporte XLSX de carga académica y docentes asignados en el período.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen de Docentes"

    institucion = ConfiguracionInstitucion.get_solo()
    _aplicar_estilos_encabezado(ws, f"Carga Académica Docente - {periodo.nombre}", institucion)

    headers = ["No.", "Código Empleado", "Nombre Completo", "Especialidad", "Teléfono", "Materias Asignadas"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx, text in enumerate(headers, 1):
        c = ws.cell(row=4, column=col_idx, value=text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    docentes = Docente.objects.all().select_related("usuario")
    row_idx = 5
    for idx, doc in enumerate(docentes, 1):
        materias_count = Materia.objects.filter(docente=doc.usuario, clase__periodo=periodo).count()
        ws.cell(row=row_idx, column=1, value=idx)
        ws.cell(row=row_idx, column=2, value=doc.codigo_empleado)
        ws.cell(row=row_idx, column=3, value=doc.usuario.get_full_name() or doc.usuario.username)
        ws.cell(row=row_idx, column=4, value=doc.especialidad or "-")
        ws.cell(row=row_idx, column=5, value=doc.telefono or "-")
        ws.cell(row=row_idx, column=6, value=materias_count)
        row_idx += 1

    _autoajustar_columnas(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_estadisticas_bloque_xlsx(periodo):
    """
    Genera el reporte de estadísticas y porcentaje de aprobación por Clase/Grado en XLSX.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estadísticas de Aprobación"

    institucion = ConfiguracionInstitucion.get_solo()
    _aplicar_estilos_encabezado(ws, f"Estadísticas de Aprobación por Grado - {periodo.nombre}", institucion)

    headers = ["No.", "Clase / Grado", "Nivel", "Alumnos Matriculados", "Alumnos Aprobados", "Alumnos Reprobados", "% Aprobación"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx, text in enumerate(headers, 1):
        c = ws.cell(row=4, column=col_idx, value=text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    clases = Clase.objects.filter(periodo=periodo).select_related("nivel")
    row_idx = 5
    for idx, c in enumerate(clases, 1):
        estudiantes_ids = EstudianteClase.objects.filter(clase=c).values_list("estudiante_id", flat=True)
        total_alumnos = len(estudiantes_ids)
        reprobrados = ResumenAcademicoEstudiante.objects.filter(
            periodo=periodo, estudiante_id__in=estudiantes_ids, materias_reprobadas__gt=0
        ).count()
        aprobados = max(0, total_alumnos - reprobrados)
        pct_aprobacion = (aprobados / total_alumnos * 100) if total_alumnos > 0 else 100.0

        ws.cell(row=row_idx, column=1, value=idx)
        ws.cell(row=row_idx, column=2, value=c.nombre)
        ws.cell(row=row_idx, column=3, value=c.nivel.nombre)
        ws.cell(row=row_idx, column=4, value=total_alumnos)
        ws.cell(row=row_idx, column=5, value=aprobados)
        ws.cell(row=row_idx, column=6, value=reprobrados)
        ws.cell(row=row_idx, column=7, value=f"{pct_aprobacion:.1f}%")
        row_idx += 1

    _autoajustar_columnas(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
