from django.test import TestCase, Client
from django.urls import reverse
import openpyxl
from datetime import date
from apps.accounts.models import User
from apps.academic_core.models import PeriodoLectivo, Nivel, Clase, Seccion
from apps.people.models import Estudiante, EstudianteClase
from apps.audit.models import AuditEvent
from apps.audit.services import registrar_evento
from apps.reporting.services import (
    generar_boletin_xlsx,
    generar_listado_estudiantes_xlsx,
    generar_resumen_docentes_xlsx,
    generar_estadisticas_bloque_xlsx
)


class Fase3ReportesAuditoriaTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(username="admin_reportes", password="password", is_superuser=True)
        self.periodo = PeriodoLectivo.objects.create(
            nombre="2026", fecha_inicio=date(2026, 1, 1), fecha_fin=date(2026, 11, 30), activo=True
        )
        self.nivel = Nivel.objects.create(nombre="Primaria")
        self.clase = Clase.objects.create(periodo=self.periodo, nivel=self.nivel, nombre="3er Grado")
        self.seccion = Seccion.objects.create(clase=self.clase, codigo="SEC-3A", nombre="Sección A")
        self.estudiante = Estudiante.objects.create(
            codigo_estudiante="EST-2026-333", primer_nombre="Roberto", primer_apellido="Gómez", genero="M"
        )
        EstudianteClase.objects.create(
            estudiante=self.estudiante, periodo=self.periodo, clase=self.clase, seccion=self.seccion
        )

    def test_xlsx_report_generators(self):
        # 1. Boletín XLSX
        buf_boletin = generar_boletin_xlsx(self.estudiante, self.periodo)
        wb_boletin = openpyxl.load_workbook(buf_boletin)
        self.assertIn("Boletín de Calificaciones", wb_boletin.sheetnames)

        # 2. Listado Estudiantes XLSX
        buf_listado = generar_listado_estudiantes_xlsx(self.clase, self.seccion)
        wb_listado = openpyxl.load_workbook(buf_listado)
        self.assertIn("Listado de Alumnos", wb_listado.sheetnames)

        # 3. Resumen Docentes XLSX
        buf_docentes = generar_resumen_docentes_xlsx(self.periodo)
        wb_docentes = openpyxl.load_workbook(buf_docentes)
        self.assertIn("Resumen de Docentes", wb_docentes.sheetnames)

        # 4. Estadísticas XLSX
        buf_est = generar_estadisticas_bloque_xlsx(self.periodo)
        wb_est = openpyxl.load_workbook(buf_est)
        self.assertIn("Estadísticas de Aprobación", wb_est.sheetnames)

    def test_audit_event_logging_and_sanitization(self):
        evento = registrar_evento(
            accion="PRUEBA_SEGURIDAD",
            usuario=self.user,
            descripcion="Prueba de auditoría",
            detalles={"password": "SuperSecret123!", "normal": "valor_publico"}
        )
        self.assertEqual(evento.detalles["password"], "[PROTEGIDO]")
        self.assertEqual(evento.detalles["normal"], "valor_publico")

    def test_report_index_and_auditoria_views(self):
        self.client.force_login(self.user)

        res_rep = self.client.get(reverse("reportes_index"))
        self.assertEqual(res_rep.status_code, 200)

        res_aud = self.client.get(reverse("auditoria_list"))
        self.assertEqual(res_aud.status_code, 200)
